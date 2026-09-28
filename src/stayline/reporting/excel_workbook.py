"""
src/stayline/reporting/excel_workbook.py
------------------------------------------
Generates the Excel workbook.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

from stayline.config import get_settings
from stayline.logging_setup import get_logger
from stayline.warehouse import query_df

logger = get_logger(__name__)

def run() -> None:
    cfg = get_settings()
    logger.info("Generating Excel workbook...")
    
    df = query_df("SELECT * FROM vw_subscriber_360")
    watchlist = query_df("SELECT * FROM vw_scored_watchlist LIMIT 500")
    
    wb = Workbook()
    
    # 1. README sheet
    ws_readme = wb.active
    ws_readme.title = "README"
    ws_readme.append(["Stayline Retention Intelligence"])
    ws_readme.append(["This workbook contains subscriber data and a risk watchlist."])
    
    # 2. Data Cleaned
    ws_data = wb.create_sheet(title="Data_Cleaned")
    for row in dataframe_to_rows(df, index=False, header=True):
        ws_data.append(row)
        
    # 3. Watchlist
    ws_watch = wb.create_sheet(title="Retention_Call_List")
    for row in dataframe_to_rows(watchlist, index=False, header=True):
        ws_watch.append(row)
        
    cfg.reports.excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cfg.reports.excel_path)
    logger.info(f"Excel workbook saved to {cfg.reports.excel_path}")
